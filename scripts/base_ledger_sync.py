"""Read BASE API v2 data into a copy of the existing Excel ledger.

Only the item list, order search and order detail endpoints are called. No raw
API response is persisted because it contains buyer personal information.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from copy import copy
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen
import json

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill


API_ORIGIN = "https://apiv2.thebase.com"
SHEET_DROP = "DROP計画"
SHEET_PRODUCTS = "商品・在庫"
SHEET_SALES = "売上・利益"
SHEET_REWARDS = "報酬分配"
SHEET_DASH = "ダッシュボード"
SHEET_ORDER = [SHEET_DROP, SHEET_PRODUCTS, SHEET_SALES, SHEET_REWARDS, SHEET_DASH]
STOCK_OFFSET = 47  # Former ③-4 row 8 -> 商品・在庫 row 55.
ITEM_HEADER = 71
LINE_COLUMN = 22  # V:AF on 売上・利益.
STALE_ORDER = "注文検索で未取得（前回値）"


class SyncError(Exception):
    pass


def valid_order_range(start: date, end: date):
    try:
        next_year = start.replace(year=start.year + 1)
    except ValueError:  # February 29
        next_year = date(start.year + 1, 3, 1)
    return start <= end < next_year


def excel_text(value):
    """Keep API supplied names as literal text, never spreadsheet formulas."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def required_int(obj, key, label):
    value = obj.get(key)
    if type(value) is not int or value < 0:
        raise SyncError(f"{label}.{key} が未取得または不正")
    return value


def required_list(obj, key, label):
    value = obj.get(key)
    if not isinstance(value, list):
        raise SyncError(f"{label}.{key} が未取得または不正")
    return value


def require_fields(obj, keys, label):
    if not isinstance(obj, dict) or any(key not in obj for key in keys):
        raise SyncError(f"{label} に必要な項目がありません")


class BaseClient:
    def __init__(self, token: str):
        if not token:
            raise SyncError("BASE_PAT 環境変数がありません")
        self.token = token

    def request(self, method: str, path: str, payload=None):
        if method not in ("GET", "POST") or not path.startswith("/api/"):
            raise SyncError("許可されない API 操作")
        body = None if payload is None else json.dumps(payload).encode("utf-8")
        request = Request(
            API_ORIGIN + path,
            data=body,
            method=method,
            headers={
                "Authorization": "Bearer " + self.token,
                "Accept": "application/json",
                **({"Content-Type": "application/json"} if body is not None else {}),
            },
        )
        try:
            with urlopen(request, timeout=30) as response:
                return json.load(response)
        except HTTPError as exc:
            raise SyncError(f"BASE API HTTP {exc.code}: {method} {path.split('?')[0]}") from None
        except (URLError, TimeoutError, ValueError) as exc:
            raise SyncError(f"BASE API 取得失敗: {method} {path.split('?')[0]}") from None

    def items(self):
        found = {}
        offset = 0
        while True:
            page = self.request("GET", "/api/items?" + urlencode({"limit": 100, "offset": offset}))
            if not isinstance(page, list):
                raise SyncError("商品一覧の形式が不正")
            for item in page:
                require_fields(item, ("id", "name", "stock", "price", "variations"), "商品")
                item_id = required_int(item, "id", "商品")
                if item_id in found:
                    raise SyncError(f"商品 ID {item_id} が複数ページに重複")
                stock = required_int(item, "stock", "商品")
                variations = required_list(item, "variations", "商品")
                if variations:
                    variation_ids = set()
                    for variation in variations:
                        require_fields(variation, ("id", "stock"), "種類")
                        variation_id = required_int(variation, "id", "種類")
                        if variation_id in variation_ids:
                            raise SyncError(f"商品 ID {item_id} の種類 ID が重複")
                        variation_ids.add(variation_id)
                        required_int(variation, "stock", "種類")
                    if sum(v["stock"] for v in variations) != stock:
                        raise SyncError(f"商品 ID {item_id} の種類別在庫と合計が不一致")
                price = item["price"]
                if price is not None:
                    required_int(item, "price", "商品")
                found[item_id] = item
            if len(page) < 100:
                return found
            offset += len(page)
            if offset > 1_000_000:
                raise SyncError("商品一覧が API のページング上限を超えました")

    def orders(self, start: date, end: date):
        if not valid_order_range(start, end):
            raise SyncError("注文期間は開始日から 1 年未満にしてください")
        summaries = {}
        page_number = 1
        while True:
            payload = {
                "ordered_from_date": start.isoformat(),
                "ordered_to_date": end.isoformat(),
                "sort": "ordered_asc",
                "limit": 200,
                "page": page_number,
            }
            page = self.request("POST", "/api/orders/search", payload)
            require_fields(page, ("orders", "pagination"), "注文検索")
            pagination = page["pagination"]
            require_fields(pagination, ("page", "total_pages", "total_count", "has_next"), "注文ページング")
            if pagination["page"] != page_number or type(pagination["has_next"]) is not bool:
                raise SyncError("注文ページングが不整合")
            for summary in required_list(page, "orders", "注文検索"):
                require_fields(summary, ("unique_key",), "注文概要")
                key = summary["unique_key"]
                if not isinstance(key, str) or not key.isalnum() or len(key) != 16:
                    raise SyncError("注文キーが不正")
                summaries[key] = summary
            if not pagination["has_next"]:
                if page_number != pagination["total_pages"] and pagination["total_pages"] != 0:
                    raise SyncError("注文ページ数が不整合")
                if len(summaries) != required_int(pagination, "total_count", "注文ページング"):
                    raise SyncError("注文件数が API の総件数と不一致。期間を分割して再取得してください")
                break
            if page_number >= 500:
                raise SyncError("注文検索が 500 ページ上限に到達。期間を分割してください")
            page_number += 1
        details = {}
        for key in summaries:
            detail = self.request("GET", "/api/orders/" + quote(key))
            require_fields(detail, ("unique_key", "status", "ordered_at", "amounts", "lines", "delivery"), "注文詳細")
            if detail["unique_key"] != key:
                raise SyncError("注文詳細と検索結果のキーが不一致")
            normalize_order(detail)  # Fail before touching the workbook.
            details[key] = detail
        return details


def normalize_order(order):
    amounts = order["amounts"]
    delivery = order["delivery"]
    require_fields(amounts, ("subtotal", "shipping_fee", "cod_fee", "total", "service_charge", "additional_charges"), "注文金額")
    require_fields(delivery, ("dispatched_at", "dispatches"), "配送")
    require_fields(amounts["service_charge"], ("collected_fee",), "BASE手数料")
    line_rows = []
    line_ids = set()
    for line in required_list(order, "lines", "注文詳細"):
        require_fields(line, ("id", "item_id", "variation_id", "title", "price", "amount", "total", "status", "shipping_fee"), "注文明細")
        line_id = required_int(line, "id", "注文明細")
        if line_id in line_ids:
            raise SyncError("同一注文内の明細 ID が重複")
        line_ids.add(line_id)
        for field in ("item_id", "price", "amount", "total", "shipping_fee"):
            required_int(line, field, "注文明細")
        if line["variation_id"] is not None:
            required_int(line, "variation_id", "注文明細")
        line_rows.append((order["unique_key"], line_id, line["item_id"], line["variation_id"], excel_text(line["title"]), excel_text(line.get("variation")), line["price"], line["amount"], line["total"], line["status"], line["shipping_fee"]))
    dispatch_ids = set()
    dispatch_rows = []
    for dispatch in required_list(delivery, "dispatches", "配送"):
        require_fields(dispatch, ("id", "dispatched_at", "order_line_ids"), "発送履歴")
        dispatch_id = required_int(dispatch, "id", "発送履歴")
        if dispatch_id in dispatch_ids:
            raise SyncError("同一注文内の発送 ID が重複")
        dispatch_ids.add(dispatch_id)
        dispatch_rows.append(f"{dispatch_id}:{dispatch['dispatched_at']}:{','.join(map(str, dispatch['order_line_ids']))}")
    additional = required_list(amounts, "additional_charges", "注文金額")
    additional_values = []
    for charge in additional:
        require_fields(charge, ("name", "collected_fee"), "追加料金")
        additional_values.append(f"{charge['name'] or ''}:{charge['collected_fee'] if charge['collected_fee'] is not None else '未取得'}")
    order_row = (
        order["unique_key"], order["ordered_at"], order["status"],
        required_int(amounts, "subtotal", "注文金額"), required_int(amounts, "shipping_fee", "注文金額"),
        required_int(amounts, "cod_fee", "注文金額"), required_int(amounts, "total", "注文金額"),
        required_int(amounts["service_charge"], "collected_fee", "BASE手数料"),
        excel_text("; ".join(additional_values)), delivery["dispatched_at"], "; ".join(dispatch_rows),
        "原価・実送料・返金等の確定入力前は利益・報酬を算出不可",
        len(dispatch_rows),
    )
    return order_row, line_rows


def copy_region(source, destination, row_offset=0, col_offset=0):
    for row in source:
        for cell in row:
            target = destination.cell(cell.row + row_offset, cell.column + col_offset)
            if cell.data_type == "f":
                target.value = Translator(cell.value, origin=cell.coordinate).translate_formula(target.coordinate)
            else:
                target.value = cell.value
            if cell.has_style:
                target._style = copy(cell._style)
            if cell.number_format:
                target.number_format = cell.number_format
            if cell.alignment:
                target.alignment = copy(cell.alignment)
    for merged in source.merged_cells.ranges:
        start = destination.cell(merged.min_row + row_offset, merged.min_col + col_offset).coordinate
        end = destination.cell(merged.max_row + row_offset, merged.max_col + col_offset).coordinate
        destination.merge_cells(f"{start}:{end}")


def style_header(sheet, row, start, end):
    for col in range(start, end + 1):
        cell = sheet.cell(row, col)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 33


def style_manual_fields(sales, rewards, row):
    for col in range(14, 19):
        sales.cell(row, col).fill = PatternFill("solid", fgColor="FFF2CC")
        sales.cell(row, col).font = Font(name="Aptos", size=10, color="1F4E78")
    rewards.cell(row, 3).fill = PatternFill("solid", fgColor="FFF2CC")
    rewards.cell(row, 3).font = Font(name="Aptos", size=10, color="1F4E78")


def dispatch_count(history):
    if not isinstance(history, str) or not history.strip():
        return 0
    ids = [part.partition(":")[0].strip() for part in history.split(";")]
    if not all(part.isdigit() for part in ids) or len(ids) != len(set(ids)):
        raise SyncError("既存発送履歴の ID が不正または重複")
    return len(ids)


def ordered_in_range(timestamp, start, end):
    if not isinstance(timestamp, str):
        return False
    try:
        instant = datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
        if instant.tzinfo is None:
            return False
        local_date = instant.astimezone(timezone(timedelta(hours=9))).date()
        return start <= local_date <= end
    except ValueError:
        return False


def sales_formula(sheet, row):
    # N:Q are manual actual costs/adjustments; blank and explicit 0 differ.
    sheet.cell(row, 19, f'=IF(AND(OR(C{row}="dispatched",C{row}="arrived"),L{row}<>"{STALE_ORDER}",ISNUMBER(M{row}),M{row}>0,COUNT(N{row}:Q{row})=4,R{row}="確定",ISNUMBER(G{row}),ISNUMBER(H{row})),G{row}-H{row}-SUM(N{row}:Q{row}),"")')
    sheet.cell(row, 20, f'=IF(ISNUMBER(S{row}),"確定","未確定")')


def reward_formula(sheet, row):
    sheet.cell(row, 2, f'=IF(ISNUMBER(\'売上・利益\'!M{row}),\'売上・利益\'!M{row},"")')
    sheet.cell(row, 4, f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(E{row})),B{row}*$M$2,"")')
    sheet.cell(row, 5, f'=IF(ISNUMBER(\'売上・利益\'!S{row}),\'売上・利益\'!S{row},"")')
    sheet.cell(row, 6, f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(D{row})),E{row}-D{row},"")')
    sheet.cell(row, 7, f'=IF(ISNUMBER(F{row}),MAX(0,-F{row}),"")')
    sheet.cell(row, 8, f'=IF(ISNUMBER(F{row}),MAX(0,F{row}),"")')
    sheet.cell(row, 9, f'=IF(ISNUMBER(H{row}),H{row}*$M$3,"")')
    sheet.cell(row, 10, f'=IF(ISNUMBER(H{row}),H{row}*$M$4,"")')


def prepare_dashboard(book, source_run=None):
    dash = book[SHEET_DASH]
    dash["H1"] = "BASE実績・精算状況"
    for row, label in [(5, "取得済み注文件数"), (6, "利益確定件数"), (7, "利益未確定件数"),
                       (8, "確定利益合計（梱包前）"), (9, "分配対象利益"),
                       (10, "赤字額"), (11, "梱包報酬額")]:
        dash.cell(row, 8, label)
    style_header(dash, 20, 8, 9)
    dash["H20"] = "取得項目"
    dash["I20"] = "状態・件数"
    state = {source_run.cell(row, 1).value: source_run.cell(row, 2).value
             for row in range(2, source_run.max_row + 1)} if source_run else {}
    labels = ["商品取得", "商品件数", "注文取得", "注文期間開始JST", "注文期間終了JST",
              "注文件数", "商品台帳未登録件数", "台帳商品未取得件数", "注記"]
    for row, label in enumerate(labels, 21):
        dash.cell(row, 8, label)
        dash.cell(row, 9, state.get(label, "未取得" if "取得" in label else None))
    dash.column_dimensions["H"].width = 32
    dash.column_dimensions["I"].width = 48
    refresh_dashboard_formulas(book)


def refresh_dashboard_formulas(book):
    dash = book[SHEET_DASH]
    sales = book[SHEET_SALES]
    last = max(2, max((r for r in range(2, sales.max_row + 1) if sales.cell(r, 1).value), default=2))
    dash["I5"] = f'=COUNTA(\'{SHEET_SALES}\'!A2:A{last})'
    dash["I6"] = f'=COUNT(\'{SHEET_SALES}\'!S2:S{last})'
    dash["I7"] = "=I5-I6"
    dash["I8"] = f'=IF(I6>0,SUM(\'{SHEET_SALES}\'!S2:S{last}),"")'
    dash["I9"] = f'=IF(I6>0,SUM(\'{SHEET_REWARDS}\'!H2:H{last}),"")'
    dash["I10"] = f'=IF(I6>0,SUM(\'{SHEET_REWARDS}\'!G2:G{last}),"")'
    dash["I11"] = f'=SUM(\'{SHEET_REWARDS}\'!D2:D{last})'


def refresh_layout_notes(book):
    book[SHEET_PRODUCTS]["A50"] = "出典：初期値は旧台帳。BASE API取得状態はダッシュボード参照。"
    book[SHEET_PRODUCTS]["A51"] = "BASE API同期時に商品IDで合計在庫を更新。未取得行は前回値を保持。"
    book[SHEET_PRODUCTS]["A52"] = "上部単価表とBASE商品IDで連携。原価未確認の粗利・在庫原価・在庫粗利は空欄。"
    book[SHEET_DASH]["I29"] = "在庫0はAPI値。未取得行は前回値を保持。"


def apply_sheet_layout(book):
    products = book[SHEET_PRODUCTS]
    sales = book[SHEET_SALES]
    rewards = book[SHEET_REWARDS]
    products.column_dimensions["B"].width = 23
    products.column_dimensions["C"].width = 37
    products.column_dimensions["G"].width = 20
    for col, width in {"A": 24, "B": 24, "C": 18, "K": 27, "L": 35,
                       "M": 18, "N": 24, "O": 24, "P": 24, "Q": 26,
                       "R": 20, "S": 24, "T": 16}.items():
        sales.column_dimensions[col].width = width
    sales.row_dimensions[1].height = 53
    for col in "ABCDEFGHIJ":
        rewards.column_dimensions[col].width = 20 if col != "A" else 24
    rewards.column_dimensions["C"].width = 22
    rewards.column_dimensions["L"].width = 40
    rewards.column_dimensions["M"].width = 17
    rewards.row_dimensions[1].height = 48


def migrate_workbook(book):
    if set(book.sheetnames) == set(SHEET_ORDER) and len(book.sheetnames) == 5:
        book._sheets = [book[name] for name in SHEET_ORDER]
        refresh_layout_notes(book)
        apply_sheet_layout(book)
        sales = book[SHEET_SALES]
        rewards = book[SHEET_REWARDS]
        for row in range(2, sales.max_row + 1):
            if sales.cell(row, 1).value:
                sales_formula(sales, row)
                reward_formula(rewards, row)
                style_manual_fields(sales, rewards, row)
        refresh_dashboard_formulas(book)
        return
    legacy = ("③-2 DROP商品（年間計画）", "③-3 商品単価・原価", "③-4 EC在庫", "③売上ロードマップ")
    if not all(name in book for name in legacy):
        raise SyncError("5シート台帳または対応する旧台帳ではありません")
    old_stock = book["③-4 EC在庫"]
    old_items = book["BASE商品・種類"] if "BASE商品・種類" in book else None
    old_orders = book["BASE注文"] if "BASE注文" in book else None
    old_lines = book["BASE注文明細"] if "BASE注文明細" in book else None
    old_run = book["BASE取得状態"] if "BASE取得状態" in book else None
    book[legacy[0]].title = SHEET_DROP
    book[legacy[1]].title = SHEET_PRODUCTS
    book[legacy[3]].title = SHEET_DASH
    products = book[SHEET_PRODUCTS]
    copy_region(old_stock, products, STOCK_OFFSET)
    products["A70"] = "BASE商品・種類別在庫（API取得）"
    if old_items:
        copy_region(old_items, products, ITEM_HEADER - 1)
    else:
        for col, name in enumerate(["BASE商品ID", "商品名", "商品コード", "商品価格（税込円）", "商品合計在庫", "種類ID", "種類名", "種類在庫", "種類コード", "台帳対応"], 1):
            products.cell(ITEM_HEADER, col, name)
    sales = book.create_sheet(SHEET_SALES)
    if old_orders:
        copy_region(old_orders, sales)
    else:
        for col, name in enumerate(["注文キー", "注文日時UTC", "注文状態", "商品小計円", "顧客送料円", "代引手数料円", "注文合計円", "BASE徴収手数料円", "追加料金内訳", "最終発送日時UTC", "発送履歴ID:日時:明細ID", "利益・報酬判定"], 1):
            sales.cell(1, col, name)
    for col, name in enumerate(["発送実績件数", "実送料円（手入力）", "返金・調整円（手入力）", "その他実費円（手入力）", "商品原価実績円（手入力）", "精算確認（確定）", "利益（梱包前）円", "利益状態"], 13):
        sales.cell(1, col, name)
    if old_lines:
        copy_region(old_lines, sales, 0, LINE_COLUMN - 1)
    else:
        for col, name in enumerate(["注文キー", "明細ID", "BASE商品ID", "種類ID", "注文時商品名", "注文時種類名", "単価円", "数量", "明細合計円", "明細状態", "顧客送料配分円"], LINE_COLUMN):
            sales.cell(1, col, name)
    for row in range(2, sales.max_row + 1):
        if sales.cell(row, 1).value:
            sales.cell(row, 13, dispatch_count(sales.cell(row, 11).value))
            sales_formula(sales, row)
    rewards = book.create_sheet(SHEET_REWARDS)
    for col, name in enumerate(["注文キー", "発送実績件数", "梱包担当（手入力）", "梱包報酬円", "利益（梱包前）円", "梱包報酬控除後円", "赤字額円", "分配対象利益円", "森本50%円", "田中50%円"], 1):
        rewards.cell(1, col, name)
    rewards["L1"] = "報酬ルール"
    for row, label, value in [(2, "発送実績1件あたり梱包報酬円", 150), (3, "森本分配率", 0.5), (4, "田中分配率", 0.5)]:
        rewards.cell(row, 12, label)
        rewards.cell(row, 13, value)
    for row in range(2, sales.max_row + 1):
        if sales.cell(row, 1).value:
            rewards.cell(row, 1, sales.cell(row, 1).value)
            reward_formula(rewards, row)
            style_manual_fields(sales, rewards, row)
    for sheet in (book[SHEET_DROP], products, sales, rewards, book[SHEET_DASH]):
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.value = (cell.value.replace("'③-2 DROP商品（年間計画）'!", "'DROP計画'!")
                                      .replace("'③-3 商品単価・原価'!", "'商品・在庫'!")
                                      .replace("③売上ロードマップ!", "'ダッシュボード'!")
                                      .replace("'③売上ロードマップ'!", "'ダッシュボード'!"))
    for name in list(book.sheetnames):
        if name not in SHEET_ORDER:
            del book[name]
    book._sheets = [book[name] for name in SHEET_ORDER]
    style_header(products, ITEM_HEADER, 1, 10)
    style_header(sales, 1, 1, 20)
    style_header(sales, 1, LINE_COLUMN, LINE_COLUMN + 10)
    style_header(rewards, 1, 1, 10)
    sales.freeze_panes = "A2"
    rewards.freeze_panes = "A2"
    apply_sheet_layout(book)
    prepare_dashboard(book, old_run)
    refresh_layout_notes(book)


def validate_workbook(book):
    if book.sheetnames != SHEET_ORDER:
        raise SyncError("台帳は指定の5シートになっていません")
    retired = ("③-2 DROP商品（年間計画）", "③-3 商品単価・原価", "③-4 EC在庫", "③売上ロードマップ")
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f" and ("#REF!" in cell.value or any(name in cell.value for name in retired)):
                    raise SyncError(f"移行後の数式参照が不正: {sheet.title}!{cell.coordinate}")


def update_workbook(source: Path, target: Path, items, orders, start: date, end: date):
    book = load_workbook(source)
    migrate_workbook(book)
    products = book[SHEET_PRODUCTS]
    sales = book[SHEET_SALES]
    rewards = book[SHEET_REWARDS]
    dashboard = book[SHEET_DASH]
    price_rows = {}
    stock_rows = {}
    for row in range(9, 44):
        key = products.cell(row, 1).value
        if isinstance(key, str) and key.startswith("BASE:"):
            item_id = int(key[5:])
            if item_id in price_rows:
                raise SyncError("商品・在庫の単価表に商品 ID の重複があります")
            price_rows[item_id] = row
    for row in range(55, 66):
        item_id = products.cell(row, 2).value
        if item_id is not None:
            item_id = int(item_id)
            if item_id in stock_rows:
                raise SyncError("商品・在庫の在庫表に商品 ID の重複があります")
            stock_rows[item_id] = row
    if set(price_rows) != set(stock_rows):
        raise SyncError("単価表と在庫表の商品 ID 対応が不一致")
    item_rows = []
    for item_id, item in sorted(items.items()):
        matched = item_id in price_rows
        variants = item["variations"] or [None]
        for variant in variants:
            item_rows.append((item_id, excel_text(item["name"]), excel_text(item.get("identifier")), item["price"], item["stock"],
                              None if variant is None else variant["id"], None if variant is None else excel_text(variant.get("name")),
                              item["stock"] if variant is None else variant["stock"],
                              None if variant is None else excel_text(variant.get("identifier")),
                              "台帳連携" if matched else "台帳未登録"))
        if matched:
            products.cell(stock_rows[item_id], 3, excel_text(item["name"]))
            products.cell(stock_rows[item_id], 4, item["stock"])
            products.cell(price_rows[item_id], 3, excel_text(item["name"]))
            products.cell(price_rows[item_id], 5).value = item["price"]
    for row in range(ITEM_HEADER + 1, products.max_row + 1):
        for col in range(1, 11):
            products.cell(row, col).value = None
    for row_number, item_row in enumerate(item_rows, ITEM_HEADER + 1):
        for col, value in enumerate(item_row, 1):
            products.cell(row_number, col, excel_text(value))

    existing_orders = {}
    manual = {}
    for row in range(2, sales.max_row + 1):
        key = sales.cell(row, 1).value
        if key:
            if key in existing_orders:
                raise SyncError("既存注文キーが重複")
            existing_orders[key] = tuple(sales.cell(row, col).value for col in range(1, 14))
            manual[key] = tuple(sales.cell(row, col).value for col in range(14, 19))
    assignees = {}
    for row in range(2, rewards.max_row + 1):
        key = rewards.cell(row, 1).value
        if key:
            if key in assignees:
                raise SyncError("既存報酬行の注文キーが重複")
            assignees[key] = rewards.cell(row, 3).value
    line_map = {}
    for row in range(2, sales.max_row + 1):
        key = sales.cell(row, LINE_COLUMN).value
        line_id = sales.cell(row, LINE_COLUMN + 1).value
        if key:
            line_key = (key, line_id)
            if line_key in line_map:
                raise SyncError("既存注文明細キーが重複")
            line_map[line_key] = tuple(sales.cell(row, col).value for col in range(LINE_COLUMN, LINE_COLUMN + 11))
    for key, detail in orders.items():
        order_row, lines = normalize_order(detail)
        existing_orders[key] = order_row
        for old_key in [entry for entry in line_map if entry[0] == key]:
            del line_map[old_key]
        for line in lines:
            line_key = (key, line[1])
            if line_key in line_map:
                raise SyncError("取得注文明細キーが重複")
            line_map[line_key] = line
    for key, previous in list(existing_orders.items()):
        if key not in orders and ordered_in_range(previous[1], start, end):
            marked = list(previous)
            marked[11] = STALE_ORDER
            existing_orders[key] = tuple(marked)
    if any(key not in existing_orders for key, _ in line_map):
        raise SyncError("注文の無い明細が含まれます")
    old_max = sales.max_row
    for row in range(2, old_max + 1):
        for col in (*range(1, 21), *range(LINE_COLUMN, LINE_COLUMN + 11)):
            sales.cell(row, col).value = None
    for row, key in enumerate(sorted(existing_orders), 2):
        for col, value in enumerate(existing_orders[key], 1):
            sales.cell(row, col, excel_text(value))
        for col, value in enumerate(manual.get(key, (None,) * 5), 14):
            sales.cell(row, col, value)
        sales_formula(sales, row)
    for row, line in enumerate((line_map[key] for key in sorted(line_map)), 2):
        for col, value in enumerate(line, LINE_COLUMN):
            sales.cell(row, col, excel_text(value))
    for row in range(2, max(rewards.max_row, len(existing_orders) + 1) + 1):
        for col in range(1, 11):
            rewards.cell(row, col).value = None
    for row, key in enumerate(sorted(existing_orders), 2):
        rewards.cell(row, 1, key)
        rewards.cell(row, 3, assignees.get(key))
        reward_formula(rewards, row)
        style_manual_fields(sales, rewards, row)
    for row, value in [(21, "完了"), (22, len(items)), (23, "完了（期間内0件）" if not orders else "完了"),
                       (24, start.isoformat()), (25, end.isoformat()), (26, len(orders)),
                       (27, len(set(items) - set(price_rows))), (28, len(set(price_rows) - set(items))),
                       (29, "在庫0はAPI値。未取得行は前回値を保持。")]:
        dashboard.cell(row, 9, value)
    refresh_dashboard_formulas(book)
    validate_workbook(book)
    atomic_save(book, target)
    return len(set(price_rows) - set(items))


def atomic_save(book, target: Path):
    book.calculation.calcMode = "auto"
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(suffix=".xlsx", prefix=".base_sync_", dir=target.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        book.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description="BASE API v2 を読み取り、既存台帳のコピーを更新")
    parser.add_argument("--workbook", required=True, type=Path, help="既存 xlsx。出力先と同じパスも可")
    parser.add_argument("--output", required=True, type=Path, help="更新後 xlsx")
    parser.add_argument("--from-date", type=date.fromisoformat, help="注文の開始日（JST、YYYY-MM-DD）")
    parser.add_argument("--to-date", type=date.fromisoformat, help="注文の終了日（JST、YYYY-MM-DD）")
    parser.add_argument("--migrate-only", action="store_true", help="PAT・APIを使わず旧台帳を5シートに移行")
    args = parser.parse_args(argv)
    try:
        if not args.workbook.is_file():
            raise SyncError("入力台帳がありません")
        if args.migrate_only:
            book = load_workbook(args.workbook)
            migrate_workbook(book)
            validate_workbook(book)
            atomic_save(book, args.output)
            print(f"5シート移行完了: {args.output}")
            return 0
        if args.from_date is None or args.to_date is None:
            raise SyncError("同期には --from-date と --to-date が必要です")
        if not valid_order_range(args.from_date, args.to_date):
            raise SyncError("注文期間は開始日から 1 年未満にしてください")
        client = BaseClient(os.environ.get("BASE_PAT", ""))
        items = client.items()
        orders = client.orders(args.from_date, args.to_date)
        missing = update_workbook(args.workbook, args.output, items, orders, args.from_date, args.to_date)
        print(f"完了: 商品 {len(items)} 件、注文 {len(orders)} 件。台帳商品未取得 {missing} 件。出力: {args.output}")
        return 0
    except (SyncError, OSError) as exc:
        print(f"中断: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
