"""Offline contract checks using only synthetic BASE responses."""

import sys
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from base_ledger_sync import BaseClient, SyncError, update_workbook  # noqa: E402


SOURCE = Path(__file__).resolve().parents[2] / "outputs/inventory-20261014/販売戦略_Life_is_Max_Bet_v3.2_確定金額・EC在庫連携.xlsx"


class FakeClient(BaseClient):
    def __init__(self, fail_detail=False, zero_orders=False):
        super().__init__("synthetic-secret")
        self.fail_detail = fail_detail
        self.zero_orders = zero_orders
        self.paths = []

    def request(self, method, path, payload=None):
        self.paths.append((method, path))
        if path.startswith("/api/items?"):
            return [{"id": 162021520, "name": "=DANGER", "identifier": None,
                     "price": 8800, "stock": 0,
                     "variations": [{"id": 17, "name": "赤", "stock": 0, "identifier": None}]}]
        if path == "/api/orders/search":
            if self.zero_orders:
                return {"orders": [], "pagination": {"page": 1, "total_pages": 0,
                                                    "total_count": 0, "has_next": False}}
            return {"orders": [{"unique_key": "ABCDEFGHIJKLMNOP"}],
                    "pagination": {"page": 1, "total_pages": 1, "total_count": 1, "has_next": False}}
        if path == "/api/orders/ABCDEFGHIJKLMNOP":
            if self.fail_detail:
                raise SyncError("模擬通信失敗")
            return {
                "unique_key": "ABCDEFGHIJKLMNOP", "ordered_at": "2026-10-01T01:00:00Z",
                "status": "dispatched", "customer": {"name": "PRIVATE_PERSON", "mail_address": "private@example.invalid"},
                "amounts": {"subtotal": 8800, "shipping_fee": 500, "cod_fee": 0,
                            "total": 9300, "service_charge": {"collected_fee": 300},
                            "additional_charges": []},
                "lines": [{"id": 7, "item_id": 162021520, "variation_id": 17,
                           "title": "Test product", "variation": "赤", "price": 8800,
                           "amount": 1, "total": 8800, "status": "dispatched", "shipping_fee": 500}],
                "delivery": {"dispatched_at": "2026-10-02T01:00:00Z",
                             "dispatches": [{"id": 8, "dispatched_at": "2026-10-02T01:00:00Z",
                                             "order_line_ids": [7]}]},
            }
        raise AssertionError((method, path))


class OfflineSyncTests(unittest.TestCase):
    def test_stock_order_privacy_and_repeat(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            client = FakeClient()
            items = client.items()
            orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
            update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            update_workbook(output, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            book = load_workbook(output)
            self.assertEqual(book["③-4 EC在庫"]["D8"].value, 0)
            self.assertEqual(book["③-4 EC在庫"]["D9"].value, 11)  # Not returned: stale, not zero.
            self.assertEqual(book["③-3 商品単価・原価"]["E33"].value, 8800)
            self.assertEqual(book["③-3 商品単価・原価"]["F33"].value, None)
            self.assertEqual(book["BASE注文"].max_row, 2)
            self.assertEqual(book["BASE注文明細"].max_row, 2)
            self.assertEqual(book["BASE注文"]["H2"].value, 300)
            self.assertEqual(book["BASE商品・種類"]["B2"].data_type, "s")
            with zipfile.ZipFile(output) as file:
                raw = b"".join(file.read(name) for name in file.namelist())
            self.assertNotIn(b"private@example.invalid", raw)
            self.assertNotIn(b"PRIVATE_PERSON", raw)
            self.assertEqual({method for method, _ in client.paths}, {"GET", "POST"})

    def test_failure_does_not_replace_output(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            output.write_bytes(b"existing")
            client = FakeClient(fail_detail=True)
            with self.assertRaises(SyncError):
                items = client.items()
                orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
                update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            self.assertEqual(output.read_bytes(), b"existing")

    def test_zero_orders_is_recorded_as_observed_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            client = FakeClient(zero_orders=True)
            items = client.items()
            orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
            update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            book = load_workbook(output)
            self.assertEqual(book["BASE注文"].max_row, 1)
            self.assertEqual(book["BASE取得状態"]["B4"].value, "完了（期間内0件）")


if __name__ == "__main__":
    unittest.main()
