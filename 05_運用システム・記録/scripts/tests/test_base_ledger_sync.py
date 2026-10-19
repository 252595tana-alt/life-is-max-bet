"""Offline contract checks using only synthetic BASE responses."""

import sys
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from base_ledger_sync import BaseClient, SyncError, migrate_workbook, update_workbook  # noqa: E402


SOURCE = Path(__file__).resolve().parents[2] / "outputs/base-ledger-current.xlsx"


class FakeClient(BaseClient):
    def __init__(self, fail_detail=False, zero_orders=False):
        super().__init__("synthetic-secret")
        self.fail_detail = fail_detail
        self.zero_orders = zero_orders
        self.paths = []

    def request(self, method, path, payload=None):
        self.paths.append((method, path))
        if path.startswith("/api/items?"):
            return [{"id": 162021520, "name": "=DANGER", "identifier": None,
                     "price": 8800, "stock": 0,
                     "variations": [{"id": 17, "name": "赤", "stock": 0, "identifier": None}]}]
        if path == "/api/orders/search":
            if self.zero_orders:
                return {"orders": [], "pagination": {"page": 1, "total_pages": 0,
                                                    "total_count": 0, "has_next": False}}
            return {"orders": [{"unique_key": "ABCDEFGHIJKLMNOP"}],
                    "pagination": {"page": 1, "total_pages": 1, "total_count": 1, "has_next": False}}
        if path == "/api/orders/ABCDEFGHIJKLMNOP":
            if self.fail_detail:
                raise SyncError("模擬通信失敗")
            return {
                "unique_key": "ABCDEFGHIJKLMNOP", "ordered_at": "2026-10-01T01:00:00Z",
                "status": "dispatched", "customer": {"name": "PRIVATE_PERSON", "mail_address": "private@example.invalid"},
                "amounts": {"subtotal": 8800, "shipping_fee": 500, "cod_fee": 0,
                            "total": 9300, "service_charge": {"collected_fee": 300},
                            "additional_charges": []},
                "lines": [{"id": 7, "item_id": 162021520, "variation_id": 17,
                           "title": "Test product", "variation": "赤", "price": 8800,
                           "amount": 1, "total": 8800, "status": "dispatched", "shipping_fee": 500}],
                "delivery": {"dispatched_at": "2026-10-02T01:00:00Z",
                             "dispatches": [{"id": 8, "dispatched_at": "2026-10-02T01:00:00Z",
                                             "order_line_ids": [7]}]},
            }
        raise AssertionError((method, path))


class OfflineSyncTests(unittest.TestCase):
    def test_existing_drop_layout_is_not_overwritten(self):
        book = load_workbook(SOURCE)
        book["DROP計画"]["C158"] = "利用者が配置した注記"
        migrate_workbook(book)
        self.assertEqual(book["DROP計画"]["C158"].value, "利用者が配置した注記")

    def test_stock_order_privacy_and_repeat(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            client = FakeClient()
            items = client.items()
            orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
            update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            interim = load_workbook(output)
            synthetic_row = next(r for r in range(2, interim["売上・利益"].max_row + 1)
                                 if interim["売上・利益"].cell(r, 1).value == "ABCDEFGHIJKLMNOP")
            reward = interim["報酬分配"]
            self.assertEqual(interim["売上・利益"].cell(synthetic_row, 13).value, 1)
            self.assertIn(f"ISNUMBER(E{synthetic_row})", reward.cell(synthetic_row, 4).value)
            self.assertIn(f"ISNUMBER(B{synthetic_row})", reward.cell(synthetic_row, 4).value)
            interim["売上・利益"].cell(synthetic_row, 14, 350)
            interim["売上・利益"].cell(synthetic_row, 15, 0)
            interim["売上・利益"].cell(synthetic_row, 16, 0)
            interim["売上・利益"].cell(synthetic_row, 17, 1200)
            interim["売上・利益"].cell(synthetic_row, 18, "確定")
            interim["報酬分配"].cell(synthetic_row, 3, "担当者A")
            interim.save(output)
            update_workbook(output, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ["DROP計画", "商品・在庫", "売上・利益", "報酬分配", "ダッシュボード"])
            self.assertEqual(book["商品・在庫"]["D55"].value, 0)
            self.assertEqual(book["商品・在庫"]["D56"].value, 11)  # Not returned: stale, not zero.
            self.assertEqual(book["商品・在庫"]["E33"].value, 8800)
            self.assertEqual(book["商品・在庫"]["F33"].value, None)
            self.assertEqual(book["ダッシュボード"]["B5"].value, 5000000)
            self.assertEqual(book["DROP計画"]["F130"].value, 40000)
            synthetic_row = next(r for r in range(2, book["売上・利益"].max_row + 1)
                                 if book["売上・利益"].cell(r, 1).value == "ABCDEFGHIJKLMNOP")
            self.assertEqual(book["売上・利益"].cell(synthetic_row, 8).value, 300)
            self.assertEqual(book["売上・利益"].cell(synthetic_row, 14).value, 350)
            self.assertEqual(book["売上・利益"].cell(synthetic_row, 17).value, 1200)
            self.assertEqual(book["報酬分配"].cell(synthetic_row, 3).value, "担当者A")
            self.assertIn(f"ISNUMBER(E{synthetic_row})", book["報酬分配"].cell(synthetic_row, 4).value)
            self.assertEqual(book["商品・在庫"]["B72"].data_type, "s")
            self.assertEqual(len([r for r in range(2, book["売上・利益"].max_row + 1)
                                  if book["売上・利益"].cell(r, 1).value == "ABCDEFGHIJKLMNOP"]), 1)
            with zipfile.ZipFile(output) as file:
                raw = b"".join(file.read(name) for name in file.namelist())
            self.assertNotIn(b"private@example.invalid", raw)
            self.assertNotIn(b"PRIVATE_PERSON", raw)
            self.assertEqual({method for method, _ in client.paths}, {"GET", "POST"})
            update_workbook(output, output, items, {}, date(2026, 10, 1), date(2026, 10, 15))
            stale = load_workbook(output)["売上・利益"]
            stale_row = next(r for r in range(2, stale.max_row + 1)
                             if stale.cell(r, 1).value == "ABCDEFGHIJKLMNOP")
            self.assertEqual(stale.cell(stale_row, 12).value, "注文検索で未取得（前回値）")
            self.assertEqual(stale.cell(stale_row, 14).value, 350)

    def test_failure_does_not_replace_output(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            output.write_bytes(b"existing")
            client = FakeClient(fail_detail=True)
            with self.assertRaises(SyncError):
                items = client.items()
                orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
                update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            self.assertEqual(output.read_bytes(), b"existing")

    def test_zero_orders_is_recorded_as_observed_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ledger.xlsx"
            client = FakeClient(zero_orders=True)
            items = client.items()
            orders = client.orders(date(2026, 10, 1), date(2026, 10, 15))
            update_workbook(SOURCE, output, items, orders, date(2026, 10, 1), date(2026, 10, 15))
            book = load_workbook(output)
            self.assertEqual(book["ダッシュボード"]["I23"].value, "完了（期間内0件）")
            self.assertEqual(book["ダッシュボード"]["I26"].value, 0)


if __name__ == "__main__":
    unittest.main()
